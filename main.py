# -*- coding: utf-8 -*-
"""
Plant Chemical Analysis - Telegram Bot
100% Mustaqil va Barcha modullarni o'z ichiga olgan yagona fayl (Standalone main.py).
Hech qanday tashqi papkalar (bot/, core/) talab qilinmaydi.
Render, Docker, Linux, Windows - istalgan muhitda to'g'ridan-to'g'ri ishlaydi.
"""

import sys
import os
import io
import asyncio
import logging

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s"
)
logger = logging.getLogger("plant_analysis_bot")
import base64
import shutil
from datetime import datetime
from typing import List, Dict, Tuple, Optional, Any

# Render (512MB RAM) xotira xatoliklarini oldini olish
os.environ["OPENBLAS_NUM_THREADS"] = "1"
os.environ["MKL_NUM_THREADS"] = "1"

if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from aiohttp import web
from aiogram import Bot, Dispatcher, Router, F
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.filters import CommandStart, Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (
    Message,
    CallbackQuery,
    BufferedInputFile,
    FSInputFile,
    ReplyKeyboardMarkup,
    KeyboardButton,
    InlineKeyboardMarkup,
    InlineKeyboardButton
)
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side

# =====================================================================
# SOZLAMALAR VA SHABLONNI TIKLASH
# =====================================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(BASE_DIR, ".env"))

BOT_TOKEN = os.getenv("BOT_TOKEN", "8939401003:AAGO2fMFsR489ZisMQO4DGZBETAt6W478ag").strip()
TEMPLATE_PATH = os.path.join(BASE_DIR, "DPPH_template.xlsx")
OUTPUT_DIR = os.path.join(BASE_DIR, "outputs")
os.makedirs(OUTPUT_DIR, exist_ok=True)

# Agar DPPH_template.xlsx fayli mavjud bo'lmasa, uni ichki base64 dan tiklash
DPPH_TEMPLATE_B64 = "UEsDBBQAAAAIABVPC11Gx01IlQAAAM0AAAAQAAAAZG9jUHJvcHMvYXBwLnhtbE3PTQvCMAwG4L9SdreZih6kDkQ9ip68zy51hbYpbYT67+0EP255ecgboi6JIia2mEXxLuRtMzLHDUDWI/o+y8qhiqHke64x3YGMsRoPpB8eA8OibdeAhTEMOMzit7Dp1C5GZ3XPlkJ3sjpRJsPiWDQ6sScfq9wcChDneiU+ixNLOZcrBf+LU8sVU57mym/8ZAW/B7oXUEsDBBQAAAAIABVPC10muw0XEAEAAHACAAARAAAAZG9jUHJvcHMvY29yZS54bWzNkkFOwzAQRa9SZZ+Mk7QRWGkWgFhRAaISiJ1lT1tDHFu2S5rb44Q0paIHYOk/32/+aKbkhnJt8clqg9ZLdLODqhtHuVlGO+8NBXB8h4q5JDiaUNxoq5gPT7sFw/gn2yJkhBSg0DPBPIMeGJuJGFWl4JRbZF7bES/4hDd7Ww8wwQFrVNh4B2mSQlQ9hkayrvXsmXV79aEb/VXCidVzPVrlfgQUE3xQL3YYKhCNzoOTk6tt26TNB18YJ4W31cPLMHksG+dZwzH8cpL6zuAyOnZ+zW/v1vdRlZGsiMlVnKZrck3JnOaL9z7rWb5TYKWF3Mh/lHhR0Hn+K/ExYFWGC6mZ86tRuOkur+WvbdDOj6v6BlBLAwQUAAAACAAVTwtdwRcQvpIGAADGIAAAEwAAAHhsL3RoZW1lL3RoZW1lMS54bWztWc1v2zYUvw/Y/yDo7kq2JX8EdQpbtvuVNEHjduiRlmmLMSUaJJXEKAoM7WmXAQO6YZcBu+0wDCuwAit22R8ToMXW/RF7kr9Em2qTNi06LA5gk9TvPf743uPji3j12klIjSPMBWFRwyxesU0DRz4bkGjUMO/1uoWaaQiJogGiLMINc4qFeW3788+uoi0Z4BAbIB+JLdQwAyknW5YlfBhG4gqb4AieDRkPkYQuH1kDjo5Bb0itkm1XrBCRyDQiFILaveGQ+NjoJSrN7YXyDoWvSIpkwKf8wE9nzEqk2MG4mPyIqfAoN44QbZgwz4Ad9/CJNA2KhIQHDdNOP6a1fdVaClGZI5uR66afudxcYDAupXJ81F8KOo7rVJpL/aWZ/k1cp9qpdCpLfSkA+T6stKjRWS15zhybAc2aGt3tartcVPAZ/eUNfNNN/hR8eYV3NvDdrreyYQY0a7obeLdVb7VV/e4KX9nAV+1m26kq+BQUUBKNN9C2Wyl7i9UuIUNGb2jhddfpVktz+AplZaJrJh/JvFgL0SHjXQCkzkWSRIacTvAQ+YDzECV9TowdMgog8CYoYgKG7ZLdtcvwnfw5aSv1KNrCKCM9G/LFxlDCxxA+JxPZMG+BVjMDefnixenj56ePfz998uT08a/zuTflbqBolJV7/dM3//zwpfH3bz++fvqtHi+y+Fe/fPXqjz/fpF4qtL579ur5s5fff/3Xz0818CZH/Sy8R0IsjDv42LjLQligZgLc5+eT6AWIKBIoAKQG2JGBArwzRVSHa2HVhPc5ZAod8Hp8qHA9CHgsiQZ4OwgV4C5jtMW4djm3k7myy4mjkX5yHmdxdxE60s3trTm4E08g5IlOpRdgheY+BW+jEY6wNJJnbIyxRuwBIYpdd4nPmWBDaTwgRgsRrUl6pC/1QjdICH6Z6giCqxXb7N43Wozq1LfxkYqEbYGoTiWmihmvo1iiUMsYhTSL3EEy0JE8mHJfMbiQ4OkRpszoDLAQOpk9PlXo3oYMo3f7Lp2GKpJLMtYhdxBjWWSbjb0AhRMtZxIFWexNMYYQRcY+k1oSTN0hSR/8gKJcd98nWJ5vW9+DDKQPkORJzHVbAjN1P07pEGGd8iYPleza5EQbHa14pIT2DsYUHaMBxsa9mzo8mzA96VsBZJUbWGebW0iN1aQfYQFlUlLXaBxLhBKyB3jEcvjsTtcSzxRFIeJ5mu+M1ZDpwCmnTaV71B8rqZTwZNPqSeyJEJ1J636AlLBK+kIfr1MenXePgczhO8jgc8tAYj+zbXqIYn3A9BAUGLp0CyKxXiTZTqlYrJUbqpt25QZrrd4JSfTW4met7HE/TtnzwQqeiy918lLKeoGTh/sPljVtFEf7GE6Sy6rmsqr5P1Y1eXv5spa5rGUua5mPVsusyhcr+5Yn1RLmvvIZEkoP5JTiHZEWPgL2/qALg2knFVq+YZoE0JxPp+BGHKVtgzP5BZHBQYAmME0xnWEk5qpHwpgwAaWTmas7Lb3icJcNZqPF4uKlJggguRqH0msxDoWanI1Wqqu3d0v1aW8ksgTcVOnZSWQmU0mUNSSq5bORKNoXxaKuYVErvomFlfEKHE4GSt6Hu86MEYQbhPQg8dNMfuHdC/d0njHVZZc0y6s7F+ZphUQm3FQSmTAM4PBYH75gX9freleXtDSqtQ/ha2szN9BI7RnHsOfKLqjx0aRhDuGfJmiGE9AnkkyF6ChqmL6cG/pdMsuEC9lGIpjB0kez9YdEYm5QEkKsZ91AoxW3Yqlqf7rk6vanZzlr3cl4OMS+zBlZdeHZTIn26XuCkw6LgfRBMDg2+jTmdxEYyq0WEwMOiJBLaw4IzwT3yopr6Wq+FZXLltUWRXQSoPmJkk3mM3jaXtLJrCNlur4qS2fC/qh7Eafu24XWkmbOAVLNzWIf7pDPsCrrWbnaXFev2W8+Jd7/QMhQq+mplfXU8s6OCywIMtNVcuxWyvXme54G61FrZerKtLdxq836hxD5bahWYyrF7OXYCZTf3uI+cpYJ0tFFdjmRRsxJw3xou03HK7lewa65nYJTduxCzW2WC03XLRc7btFut0qPwCgyCIvubO4u/LNPp/NL+3R84+I+XJTaV3wWWiytg61UOL24L5byL+4NApZ5WCl16+V6q1Kol5vdgtNu1Qp1r9IqtCtetd1te26t3n1kGkcp2GmWPafSqRUqRc8rOBU7oV+rF6pOqdR0qs1ax2k+mtsaVr74XZg35bX9L1BLAwQUAAAACAAVTwtdv/1vDlcLAABjYgAAGAAAAHhsL3dvcmtzaGVldHMvc2hlZXQxLnhtbI3dXVMiyRIG4L9icHXOxRnoT2hDjRhBLKpq6YmdOOdco7ZKDB8OMOvOv98GutVu833Ji41Vn84mSemayqILL17Xmx/b56LYnf29XKy2l53n3e7lvNvd3j8Xy9n2y/qlWJXyuN4sZ7vy281Td/uyKWYPh6Dlohv2eml3OZuvOlcXh59921xdrH/tFvNV8W1ztv21XM42v6+Lxfr1shN06h/8OX963u1/0L26eJk9Fd+L3X9fvm3K77pvZ3mYL4vVdr5enW2Kx8vO1+A8D5JsH3E45H/z4nX74euz3ezue7Eo7nfFw+GhdusXXzzuhsVicdn5o/zB/snerdc/9kdPymN6+5wPEfsHmZX/+6s4Hv11nJWp/jw+7vjwmN23B/r4dZ3A+FCg8gnfzbbFcL34//xh93zZGXTOHorH2a/F7s/1qymqJ528PYfRbDe7utisX882h5QPPPjSTzpn97+2u/XyLaZM9X5/0Nd9DY9P77IzX+3L/H23KXVeprK7Wq2X84vurkxx/333vvyvPPvbQ4SqhwgPDxGChxjNfs5/zpoPcoi7PhH37ZsRooY8KkzOlj8WQtzoY9z+1fcmN/yMSQ+ccQzPeMvP2Ec5GnjGCT9j0ENJWumUjd9xpPodR63T/HXVu+j+9bG+xyOiz7mPoRgoVpJG2rEq7fhT2kkr7RimDcVAsZI00k5UaSef0g7a5U5g3lAMFAvFtVN5Ew9jpgl9sS6Fl2kuPUyjcKmqcOnnwrV/4SksHBQDxUJx7VTeC3eUWChcSgt3JxVOephG4fqqwvU/FS5sv+L6sHBQDBQLxbVTeS8cjJn2aeF+S4X7/IyT92fcqOBAVcHB5wq2X3oDWEEoBoqF4tqpvFfwKIlQwQGt4N+X//r9n7t/d8Vr9xia4pdgpipg9qmAUfsleDykLxQQioFiobh2Ku8FzE4N8IFuhA+SD/My+US6ES+A48w1piGmEaYbTGNMt5gMpgkmK1KzdrpBL4BDzSjAox0mg8mK1ExaN84E8OoeBXiAwWQwWZGaSeuu7QBeUaNAuqSqpCEZTBaTw2l4HDXFUblIzRaqp+uherBAFUkFwmQwWUwOp+ErkqYtOCoXqVkgXR8bBrhAAS4QJIPJYnI4DY+jpjgqF6lZIGUXDpvHUUVigSAZTBaTw2n4iqTZB47KKyLTi1DXwobtHvZDgSI4scBkMFlMDqfhw5Ndb6QbTaLeqdlFpLvqIvjCvcY0xDTCdINpjOkWk8E0wWRFatZOd0FG+IKM8AWJyWCyIjWTVq7z4IskIis8ZImHrPGcfrnrVnmi9jLPh6Tx+g4mg8licjgNj6OmOCoXqVkgXbcRwVWcUYQXkjAZTBaTw2n4iqTZBY7KRWoWSNdFRbgfivCCESaDyWJyOA2Po6Y4KhepWSBdqxThVinCrRImg8licjgNX5E0u8BReUVkdhHp2rIIt2UVSbMLTAaTxeRwGj462cxFWeftIY6NzQAfHOvqEg9OTUViXQsZwy7sGtMQ0wjTDaYxpltMBtMEkxWp+XaCbj6Y4O4ywd0lJoPJitRMWjf3TPAsMsEdHyaDyYrUTFo36UvwpC/Bkz5MBpPF5HAaHkdNcVQuUrNAuglmgieYCZ5gYjKYLCaH0/AVSVMRHJWL1CyQbjKb4MlsgiezmAwmi8nhNDyOmuKoXKRmgZRvjuLJbELeFSVvi5L3Rckbo+Sd0QRORXBUXhGZiiS6yWyCJ7MVSVMRTAaTxeRwGj6RJrPNt3t1w20anppdpLphKYVX9jWmIaYRphtMY0y3mAymCSYrUrN2uhErxSNWikcsTAaTFamZtG4USfEokuJRBJPBZEVqJq28vQFf2Sm5r4Hc2EDubCC3NpB7G3CbiqNykZoF0rWpKW5TU9ymYjKYLCaH0/AViXd/4DZVpGaBdO1YitvUFL97iMlgspgcTsPjqCmOykVqFkjXZqa4YUzxO5WYDCaLyeE0fEXS7AJH5RWR2UVf10v2cS9ZkTS7wGQwWUwOp+H7JzvQvm647aenZhd95d1V8Mq+xjTENMJ0g2mM6RaTwTTBZEVq1k43YvXxiNXHIxYmg8mK1ExaN4r08SjSx6MIJoPJitS8gU13ZQ/wlT3Aq0SYDCaLyeE0PI6a4qhcpGaBdCtSA7wiNcArUpgMJovJ4TR8RdLsAkflIjULpGvHBnj1a4BXvzAZTBaTw2l4HDXFUblIzQLp2swBbhgHePULk8FkMTmchq9IvEUUr35VRGYXA10vOcC9ZEXS7AKTwWQxOZyGH5zsQDPdcJudvEkj0w1LGb5JA9MQ0wjTDaYxpltMBtMEkxWpWTvdiJXhESvDIxYmg8mK1ExaN4pkeBTJ8CiCyWCyIjWT1l3ZGb6yM7xKhMlgspgcTsPjqCmOykVqFki3IpXhFakMr0hhMpgsJofT8Bm+SQNH5SI1C6RrxzK8+pXh1S9MBpPF5HAaHkdNcVQuUrNAujYzww1jhle/MBlMFpPDafgM36SBo/Ls5E0ama6XzHAvmeGbNDAZTBaTw2n47PQd9z3lxpDeydWLoKfcJ9HD6xfEhsRGxG6IjYndEjPEJsSsbK0yKndu9MjWjR7Zu4HNELOytVJX7t/okQ0cPbKDA5shZmVrph7optmH41DqAV7XIGaIWWKO5OJJ3JTE5bK1SqX8rIAA9wS1yaXCSxzELDFHcvG1SRMREpfL1iqVrm84HAdLhTsHYoaYJeZILp7ETUlcLlurVLpu5XAcLBXuV4gZYpaYI7n42qTJCYnLayPTkyDQ9UiH42Cp8PoHMUPMEnMkF18bG56V+7qC8OMnlIBTKa9JvIvomtiQ2IjYDbExsVtihtiEmJWtVUbl9Uo2SgXiFqX6KZPrFZuVrZW68voJyfUT4lUGYoaYla2VunJDeIjb/9rk1PECADFLzJFcPImbkrhctlaplJ1NSHavh3glgJghZok5kouvTZym4LhctlaplJ1bSDqwkOxxx2aIWWKO5OJJ3JTE5bK1SqXszkLSnYWkO8NmiFlijuTiaxOnKTgur41NU0JlNxiSbjDEH9JBzBCzxBzJxdfGhmfl/qeg2ijEpinKnUIB3qNzTWxIbETshtiY2C0xQ2xCzMrWKqPyRRiRF2FEliSwGWJWtmbqsXJJIiZLEjFZksBmiFnZWqkrp+sxWSKIyRIBNkPMEnMkF0/ipiQul61VKmU7EpO2IiZLBNgMMUvMkVx8beI0BcflsrVKpWw5YtJyxKTlwGaIWWKO5OJJ3JTE5bK1SqVscWLS4sSkxcFmiFlijuTiaxOnKTgur41NU2JlSxWTlqoycZqCzRCzxBzJxdfGhmflnskgOb2aotzJGODdgNfEhsRGxG6IjYndEjPEJsSsbK0yKoc2suExIDseiRliVrZW6sqhhmxFDMheRGKGmJWtlbry0iebBAOyS5CYIWaJOZKLJ3FTEpfL1iqVcjWFbBcMxI16danIago2S8yRXHxt4jQFx+WytUqlbAITspqSkNUUbIaYJeZILp7ETUlcLlurVMomNyHNakJWU7AZYpaYI7n42sRpCo7La2PTlETZyCakkU3Iago2Q8wScyQXXxsbnvun/mU5Ht798Gn2y2LzdPgY/e3Z/frXan/c/sC3H799mr89fPr/p5+Hewll66Wl9VLJ0vDcpqEkUe/cRj3xfEl5ukSSeHBu44Ecs88vkfOL+qVFfcmyMotMzKJfPqn+4Tl132t3/LMIf8w2T/PV9mxRPJZ17H0pf8WbY+0PX+/WL4evyt/M3XpX/mbq756L2UOx2X9X/nIf1+td/U35IA+b2et89XT8gw/nG82ffFg/Ps7vi9H6/teyWO2Of/NhUyxm+7+bsH2ev2zLtM7nD5edzeThUJju29+VuPoHUEsDBBQAAAAIABVPC13N4o62ZQEAAJ0FAAAYAAAAeGwvZHJhd2luZ3MvZHJhd2luZzEueG1s7ZTRboMgFEB/xfgBQ7DVzliSpWbLXrb+AkGsJALmQms/f6B9UJctS7PsaU9w74Wj93piOdgKoqvqtN3HrXN9gZDlrVDMPpheaF9pDCjmfAgnVAMbpD6pDpEkyZDtQbDatkK4aqrEtHSDOYiue9K8NUDLBoyiJTcdxdsShTUE701D88xDxlSISjAD9WFYwj7k8Ibs8jE1nkATy5mJR8iSl+5wsgZisiR6YEZmxMA6AetbyZ+BKUFLfXmZxUffAX+7HCGS9T7GcaR9ch8fWgYuwjEaq8sLPoc+Q64NKF9hxe1p08wLdsfUFZM6nqEq5lh0BnkHioc+PIsX4+72Uvxu0g0APwGYppFcVIafldBuooDomJNG21b2No6gCFOH13qcNFp0PI/9fvkReSc9MxwL91ZGfikoSZdCZVmWb78XNHncpl8bmiYrQ3O8WRPXhmKvMb5fUbJUlPwr+keKkt9UFIXfMv0AUEsDBBQAAAAIABVPC10wuwbIlAAAAHYBAAAjAAAAeGwvZHJhd2luZ3MvX3JlbHMvZHJhd2luZzEueG1sLnJlbHPFkD0Kg0AQha+yzAEctUgR1CqNbfACwzrrStwfdidgbh9DUiikSJfyvQcfH6+58kIyB5/tHLNa3eJzC1YknhGztuwoFyGy3xYTkiPZYpowkr7RxFiX5QnTngFds2eq4RH5F2IwZtZ8Cfru2MsXMGpLSUANlCaWFnBd3tVnqYqNCKofW0j9WAH+S6Q+iNQvETzc3D0BUEsDBBQAAAAIABVPC1139yKFjQAAAPMAAAAjAAAAeGwvd29ya3NoZWV0cy9fcmVscy9zaGVldDEueG1sLnJlbHONzzsKwzAQBNCriD2A106RIsiu0rgNvsAirz7E+iDJxLl93CQ4kCLlzMCDkTdeqLoYinWpiM0vofRga00XxKIseypNTBz2Rcfsqe4xG0yk7mQYT217xnw0YJBHU0zPxP+IUWun+BrV6jnUHzDOmR4uGBATZcO1B9yWd/lZu2Z3QYxzD3mcO8BB4tfH4QVQSwMEFAAAAAgAFU8LXe0sgBiwBQAA9DIAABQAAAB4bC9jaGFydHMvY2hhcnQxLnhtbO2bXXPaOBSG/4rX07ud1N98ZICZhLbZzKZbpnR7sXfCFuCNLHllkUB//R5JNhhis6UBtsk4FybWsY6kV0fPMZHSC+eIi3GKQmwsE0KzvjkXIr20rCyc4wRlb1mKKVimjCdIwC2fWRFHjzGdJcRybbtlKRfmoKc+Bz0RC4LhYzno8TicD3rocsKi1YjrBi7RDzSRoJialnSVPteLdMLlBbpqFVeurilcdZct2Xv2gDlBK+MBkb5py+azdKSqUvYhJuQo4yH0GAPSHdIC8Uy8Q9lcdztjJI5kS5ZsCj70EMRSD+Q4E2NwJqRCRpaGH2Jo/w5lYoQ4OHRMA1QUn+AyJeyxb2JC4jSLM10OLTL+zTQeOUqhs/8sEMemgWgIxX0zFLy4GQounWnJMjEWK4J/nnjKwyLC08+gZvYNuurboMdEqRJrbQSP73HfpGysfjONe8wpPOnaWjn11ARlmMQUyxvpUs2fnFt1I3uGh4TruRVLRz1DFslHFumyVmDbdi7TIvk0nepiryi2Sl70XbkBgkRMDbFK8RSA0Dd/TegFEdodRjsGjLQhzHYMYZa3pPXQS0srVLXwwDWNZLRI7QiiMwgSevHnOHeSqtWo3ORkSQkTVxyjQQ9WJ1sIuTBDJATmQ02g/E4HiV4GCWNi/hFx0Fz6fUB8NWSE8ay8ujE0EkfLUhHjEeale4kFmMjPeDroTQfjOcbC+eXN9Ru3Z02VZYhAXeiiGLIFFbqmCttUGOBaT+vD4N1o9FvPeoAxpVKHTUWr8K4IVPDm+ZAwYOk5XTuAEAvlUuM0+s/4QmGIqXCKsCk/WQMZmE8YdvQEN4kSHsazSiaM6EphzEOCle7xt3yagjJk1w0eY4V+/9COo3U3cIMTCLwtq1XouvyKoBZdJDuBefXGu4RLV0UnmPMg0wMYsggPbjDFHBF4YFO2Hb3tp9Frb0K3MDjKEDw1uMrgVFTxtKWijq8sbkWdQFsq6rSUxSvVsTYDtgppLK3Uqkava6nX9fH1ett1vTrNaozuPqO3z+jvMwb7jK0qY6WMWkBN1RLkLIXQbY46Oxx16jk6PJijbmAk9+SFkNR9vSStHNrZSXqIwA1JT0rSodRreG6Sdlq1uoExqDB6a6NTT9KO364nacf360na8Z3jkdTdIalbT9L3B5M0sF8QSb3XS9LKoZ2dpIcI3JD0pCR9L/V6f36StvaR1N1D0ioerknq7SOpV+F2TVK3fTySejsk9epJenMwSdsv6Z3Uf70krRza2Ul6iMANSU9K0hup183ZSeruI6ndtks/Fe+gBVbb3Qo3BVbbnT1YbXfqqdpu+8ejqr9DVb+eqrcHU9WxX9ILavB6sVo5tLNj9RCBG6yeFKu3Uq/bs2PVrnjNdPcgbw3SVgUr1yANKii7Bqm/j6TPAml0NyEZWOfs8Q7PMI1+x1v7zlAO1XdKhkj8gRK8UzrGvKJ0hLlcNjul14vJhODxGiV6izDvDFre5puKvhMEXd9xOwpX28WB127pcWxtwcEDV8taH/AsAQDIbBFDr5CIGdWPJTH9iJaFQ/1QhAkW5R6i5YjlG3cTeZ+gvxm/4XEkd06zY2cARSWdAKYECfg1SSPAC52ZBiIzqvepf2Sz1qnerO1832btgenE2lUJIvNDIozNmuyb+UI1jYwtIF7uYnqPozxQVfUvcXgvd1F1m5RRlYtg0upMAkohnNbzRfFSfGE/7VmK0863Wz3f7ZPM9ymOeFy08hfV5qjH5qhH97CTHs0Jj/oTHiFnWXa1rEgvyoJziKCFYH9hztaGayweMV6nkAgyo/KdZ6HqZFS4Pl4yIk0yapJRk4yaZNQko1eZjIqvL89KRlbp+KL6qld8bn+naQ4fN4ePGwj87xCwijUqV+3XOPtEyar015soztJrqH+fXeWLd4ZS5ST/Z4hz/F12Mjv1QdZX8MK5+T+Xwb9QSwMEFAAAAAgAFU8LXa7AGjedBAAALx4AABQAAAB4bC9jaGFydHMvY2hhcnQyLnhtbO1Z33PiNhD+V1zPvXWCbRJykAFmEq656RxpMiG9h74JewEVWXIkkZj89V1JNjGcuUsu0Ek75MGgXWl/fNLqW5xuPCNSjzISg5enjKueP9M6OwsCFc8gJaohMuComQiZEo1DOQ0SSR4pn6YsaIbhaWBN+P2u/ex3NdUM8CPvdyWNZ/0uORuLZHkjnYMz8hMuUkK5HxhT2VutGCPSPDDUoHxK+8zw6UIOTPTiASQjS++BsJ4fGvcqu7FLubikjO0kH8Z3kZALyAEklf5E1MyFrQSjifEUGFf44VLQuUtkNxvjSaENQp7K4kuK/odE6Rsi0WDke4iivsbHhInHng+M0UxR5eToUcgn33uUJMNg7xdEgu8RHqO458daloOBlsaYg0zpkV4yeD/nqTgWCUxuEU31hKGehIjH2KJCHTZa0jn0fC5G9pvvzUFynNkMHXJ21pgoYJSDGRiTdv/M3tqBiQwGTLq91Xlk57BFeiUSJztthWFYwLRIrycTJz4uxUHFihtVHTCiKff0MoMJXgg9/9eUHzHtzAHZUABxilhtKGJVeHJ4uNJyCNUVHprmiTktBjtG+BQPCT/6c1QYyWw1WjPFzZIxoc8lkH4Xq1MstCnMmGgNcuBuoGLkDokrg1QIPbsiEjE3dh+IXA4EE1JVqxvQCU3yikjIBGTNBfD2qvWwFqJO2MI9j83Zlzz54YaTOAauo3IfqzO3VD0CLBY8+ab+U4sE5rNMx4K5RTGVMQObJX0qcGtVk1453EXJvDy13WDdaTVbewB4HdagxFVLPNOmjv87R2apPgldwxRlJndY426qGRFbRglV2e3ovvAVlaLf7iuClYHhGN3yRXqZas+lOhAJXhmfgYMkDO9AsZAxDCmfQ3Ig3APhvopwO6/j2wPPfo9n10o2qFxm+VfiivgWJv3upD+aAejolw9fPrTO8NHpBhOrHhBMHiesyrxfVDlOeJZ1M/xccF3hmkx7yMBuVx76YTd4wAAy/ayIrKLZ+lbTtJpWzZpjq/lYs+bEaqKwsih4Dj8oEw1c3sst2Q9N9sPvZR82jIudpN5utNth5S/ahkQUNk7qsVgPbD0uXHfaiGoWFlB1Gu3jH4DlYHItV4UIAttfJXimFGpn4nEIUzxYX2DtFxbKcfmGZED0HySFDekIZI30BqShvA3pxWI8ZjBaNTauSItgSP57UdYnUbMVYQ23bcmtizut9onLY63ZxAnn+VYbOBcrZ2o6SYpRYd0L7qallF+RvDToJiXAQFcjJPmNKFrUsRmn5G8hP0tqy1HturmwPZLrLSZ4ReHXNEMmVnyKrMCm3BHEz1yXUf112X7ZdfnK5jbYROnlTUe0AvmOxnPze8H55ILbzhg3bZtKoxSP02q/OOT6TrzbJma/+92s3++Pe9nvffRWR6fF9XrosQ491h56rFgKpc7zGnqxGiguEbLQ4i+QYqW4AP0IsKKQBJnR2i5YqJ6MStO7IyN2IKMdkpHpDg9MdGCiAxMdmOgdMFH52+VNTBQ8v6U3375Sdc3ZcuNF4QUGNlfnhYspyayF4v+I/8Zr5/F036+c/wcM9vwv4v4/UEsDBBQAAAAIABVPC12+CDD1iRQAAIyNAAAiAAAAeGwvZXh0ZXJuYWxMaW5rcy9leHRlcm5hbExpbmsxLnhtbJWd3XbbOJLHX4UnN717TiwTxBfZX+fYDofNlTvptNN7pueOSZSYbVlyKNkT53qfYR9wn2QhmwBYbBRUmIueSPUTxfoLBP8oAvCPq6/71bDp1pf95ib7erve7H56cb3f331/err7cL267XaL7d1qYyKftsNttzcvh8+nu7th1X3cXa9W+9v1aZHn6vS26zcvfv7RHu98ux2P9/1AOeL206f+w+rV9sP97Wqzfz7ksFp3+3672V33d7sX2fB9//GnF0P7kZnvefru193tajf5d/bQrX96wbLd/UPG8ux9N7w4/Vu4iIfNp+/6YR/9fBxo1x/7b4H3m2E1PH4a7vt9ttl+7IcuwLw2qpmjfu4Dsbf9h+t+/6Xf7gPBy/52uzm57B5v+QkLxJf9Q3/ya7f52A39Jgszv5pP32/+Msf5tVv3my4By4o8+2qEfbxfBz5x1e12/ZfsZvswdDcGHcl9P/Trz90m8IlX3br7cL19JMG/7brhcCaB0E037Le765vuWdXDj/Y1gPX7fX/TbbqbQGx/f7Pus9XDustk6LPL/rG/ybbf7cPh5n69fexvA5F33Z1R7/b9fagRvep2+254b879cNKhA7/58qUzLW24NcqGibO71Xpnfu0mY1noh5zGeSDOzfvdXfb5UxYImi/vbs2FYJqC+aG+HH6o0Emubzu0Qe+ut/v7PRpedjf7+10g8Ge3vuu/IUl32SZjP4TSPUT4D0Elnj8U0qDLsm/o8Z5C4U8dvir0oc/D6jEznUDoGv7j2/0tqsb//c//jl1TOFjEgjwWFGjw4rrf7bf7Veiabne33dr87thnr7r33/rxqsMY869ub3oELP52O3TZ2V/9kF1Fv+WZ2xxA09LpME+BzTWCwlYMm1DsLOZs7CT+flz8HN6a/8t+y0iM+U4072tzkzBJG/DX7mb9eD/03VM+dDZ0RQRRHrwYm+13u+v94S4dCaK/2ofroTM3SbzNmTvs/XD/+bvQRfbUXT1HI18xmB73ifwjTr7tv/VrEnk4qTuT2T543z+0wYvr1W6zxS+VOqTkl8MnP9hPBvX8735/aFv9B/TIu393m912g8bPdtnVdf9piza6s6urX9p/vEFO4Ozq3ZurX7DgLntnboHooa1NOFl269Xm4725R7MT9tIYkpfsJNwM7x/M9fTHl2F7l/1Xt3lvfqCiMHRRvCxYEbpLD/d3h7M7wfI3tu7g6vYZqtD5sP3yZZu97YbtdYe33bPDjdiceTaefvacNnKfcw7pbLjrjE/oH7chA3S+HbY32+zQEr6a7zaH/QGxdcuTw3eb/4pQcDBm5WGbPa6G7N40UwOO7/SI2704OTkc8PB/xcvs6VXoB/lz9Vf3sDL+7NlUm+4h3BJunr6uO9l1u83zj3zI5SWSy+9nf/7zzevsP67OXp2dvHvTfPefJ6/O/vH6LNgitt+Zq+QpjVFwcxJX4eZo7gPm2rciHMCnz4Wu2j82H0fHmnXmFwoQTzfHdTc8jUceOrxdNCfLf55nlcm5NFrqoI7OlGe33cfF4zZ0+u/q38+WWXv5qv1Xm717c5n99ubqXRv6xs12MP3Rbr/eBr3IJDcj33rs19Yh9NO9uSxvXmbf+sH89+N2bcYwoXb6Gjc3r84uzy5+efPnSXP2r7PX5getf//1zesDePq3UeArc5u8Wu0nr7Knf7Vm5Ji7Y8P3GfJ+gbzPkfcF8r5E3lfI+xp5v0Ter7C80ISxjBmWMsNyZljSDMuaYWkzLG+GJc6wzAss8wL9rbHMCyzzAsu8wDIvsMwLLPMCy7zAMudY5hzLnKPNHMucY5lzLHOOZc6xzDmWOccyF1jmAstcYJkL9ArHMhdY5gLLXGCZCyxzgWUuscwllrnEMpdY5hLt3LDMJZa5xDKXWOYSy1xhmSssc4VlrrDMFZa5Qvt1LHOFZa6wzBWWucYy11jmGstcY5lrLHONZa7RWxqWucYy11jmJZZ5iWVeYpmXWOYllnmJZV5imZfo3RzLvMQyr7DMKyzzCsu8wjKvsMwrLPMKy7zCMq9QI4M7GdTK5IeHAcP239lwMGw///hhtV4f/n1evMj2ZvC6H8ybDz+/+u23X348fTCW8QB47AJihcxub9YhsIagzDGwgaBGj9hC0OQ4J09NXi45PknujD99JA8d9nyMLSrO8+n/ZDD/JLpOopskuiXRQBIxlUQ8/y5BSUSSJJYuK0gXQUkcrVlQBBfnZTX9nw6K4OicKoKciiDHphRUQSapYOlSCXDeVVAFR/McwCqoiaX1TGER1MTRMqjwUuIXw2UgBtRTU/XUs3rhNqSS1FNOvZygnqMLSVDP0roU4Ex4UD1HC8KZLNXYDQbFHINsARsvEFRPBdXPHwg3R50kqKVLCZtjGRTU0cFrtrFhreXxK7x1dFDgpR5vCEHJRgXUguOSlVPJysgPcF4mSVY6yXKCZI5mUOBw125prSqCgo4uiuNnsizHG2dQ0OegqBYFqmc11bN64nm4CVZJelYuC03Q09Kq1ODY4Wva0kLz48duLc0FO94DLCvrL4KCPkcVXwj8HsOmV/W5eUWyVjMu4q1mZMRczciIu5qRx+wVA1chKyMGi7nLsCDcvi7S8DoNb9LwloZDZcD1xKqIz2JVmjIWp3mMCS6P443HKXfV1uM5vK3il0WRT5V5eoWarzFKlcbhtPvdBGfH8cbjucZzddI4XIc7bBsP9i+hIFSRARVZzISNUbKKzMlCuQV6XMBeOOjVG4frShEuPY8rKHrwnrkcccSKjdFyUZbgWCUucwFkLmLWbIySZS6cboois8O5oshcuNYHdQv7XY9L+KuEBw8jjti3McrMjRkXFozMCx4zcAVPE5anNMja4xz2AmHP4fDwZd36uECuex5zamOUqQWLiAfG8IWIubVCpInnxtKzkQ8insMLRRHP4lrD20XYAXs8eLLLMY6ZNBuuzNmhFxgsF0171XPzimTZZlzEss3IiGWbkRHLNiOPVsRAj8aLiGXjvkMrCE0nDa/T8CYNb2k4VAYWC3nEsnGepozrYiQ0VeGu3uOU6lczwaHrCVs2jzM4IsMtGwf9DbdXXVgakSaN60AoHqz2OKmz9jjJzXo8h3fBYCtYWjzY+YSCUFRQg+Qy5uC4TBPVFRYFZdDtcVIn7nEGmydSofZ1S3383Jcjjji4MVqauy/4aobLDIqVXMUcHFdpMrsKpIAXXvhe6fEC+hLkOYDDGaz/IB2eK1lWjCKzijm4McrYoojcLEHVkuuYheM6TVlXihQwd0RZX7mEpgVR1uE5HAcgDdiVMktY+AwXjkYc8XdjlBWLiLCgqsLLmL3jZZqwrmBJsncej90rvLAOj1UDvLCuwllSKnIjjrk9G5aLokKlFdPe4Ny8Ihm8GRcxeDMyYvBmZMTgzchjBk+Ay1LoiMET7qpkxfEf4CINr9PwJg1vaThUBlxXoowYPFGmKeMafvh5Zu0BSbgjNx6f1f7DF5LDdQWvO9zsClCfFLbaHRajShMjrUDpcQHv6Yg0vuIIH6Ig0rjnD5SnVkuLB3uYUBA+VQalTZnHLJ3Mk0R1eBk0ZbWPU+6YjcNNa8RbjFPR4wqeLPJsOY95uDGqFzn+NEqC8qZkMdcmWZqSvrxJKRJ7nEuKsK5eWVbHT6b1uAj+rssRQHzaGK0WPCIlGPDLImbTZJEmpatJkgZvHicN3qQvYXKKlA7nsIsNP7MfccSmjVFjJmKNFBQMJI/5NMnTlPU1TMoIzuOkEZzDaY/yPV4QeujliGM+bQwX+SLHn+6racXh3Lwi+bQZF/FpMzLi02ZkxKfNyGM+TYHxv4rMxzlXSROFLtLwOg1v0vCWhkNl4PwiFfFpSqUp4wbJOeE6qR2uwv1l4wApKA+zHC4ELIXg5QsF3LzSMZ+mdJoYFleKYucdLiXBeDUOFyLcqziAK1gOQeZYRQYyl6EglBFYf2WHimEZyzQZLa444dFZ7XBzyzze9zYO57OyTvDorcOLSlFEjU2curTRytyFcGXBQEJVMaemqjRl3WwgRmghtcNFSci9cTiXhPlVrcMLZELlCCBObYwKvVB4nV2D8YOOuehznSdp6XBZUZ42O1xISit1OGfhGboOKIpwPXIEEDs2RpVYwJleESnBAEKzmDfTLE1Ki0v4WBG5iThc8HBP6YCiojREh7OqPN4jLEccc2NjWKsFi0xPnfae5+YVyY3NuIgbm5ERNzYjI25sRh5zYxp0YzpSBjjXrhejjHou0vA6DW/S8JaGw4m2oFMq84gbK5MmVF04XDPC0Lp2uOLyeMNvHC4LylwDh4sCnjteTC5BH1PaCy0sDUuTxuKqIhQAa4fLilKkcLhQsAsJj6wdzhVlZG3xYH8TCkJRQbmitOPwsKhJ9YcLh6uS8Hiidjipn28cLgTBB7YO55TFDssyVre5tFERcxclKFeUdhgeVpanKWtxpQjlxNrhUhIewDUOFzll7YbD+WzGU/gGOeKIbxujIo/NxyrB/IhSxHxbmTTh4cLhajYxLzxIc7jkhGeVjcN5eIZV64CC9IhuxLFVCONqs2o2nwtqCcoipYwZt1KmaWlxJeCAN7xGzeFhw9+4OBeU+qTDi1mFHrng7ao0pF0+hzVbCPzWXYG1nRVxcWdFXt1ZkZd3VsfWd8LzBj1VFVu3WbmOivKw6SINr9PwJg1vaThUBvQ0VWz5ZiXSlPFzn2DPEb45Vn4mJ+zCwwNEj8tI5csr4/BYnQ4qA/qNKrqms0papnnhcK3V8Vxrj3NKHcLhqoS9UviBpseRdUuxhZ2hIFQRlF6r6NrOSqWp6GqvilDJqj2eU9yXw5WCzTd8X3O4rAhzxJdVdIWnk2nBI30xKORW0UWelU5T1k1QQq5VV+gNjyF9XBDWbrcOl4ow/29ZRZd6WiHEgkd6PVC9raKrPasyTTs3dYKy0rd2uFKEyYyNx4uwKXCAFJGZG17L6CrPytdr8UkwFVzoGV/pmTST48LhmlMqHQ5XkmDrG4/nhApm63DJKCtaqiPrPSs7DOD4AIvlYHbZ4SVtyWdOnl82R2OLPnPyDLM5enTZZw46s+eXmENz0UVJeYRxkcjXiXyTyLdEfiYQXBibx+aa2ShdoKQHVLXnSatumkS+9bzS1OlnLIfrY/PoBDQbpivkeAb58A3S8+FhcHMMaD2gZpPiw6M7xwd7m2AU6sdAVfb5JerUbJisH/PTwwjPtesJzyizzxL51vNSE8qhS8sjfs2pxRcSN2yMMShxdBKaDdMl9hPFKIUdz6uKstVEIt96XgrK+gHLI7bOqSUWFT6BirECShydnGbDdIndjDBNWDtce15pygyHRL71PM2SWB5xe06tIxJzKHF0lpoN0yX2M8kQTZMqw00i33peIEM3S2DOzsaFWBQRGcFC8PPDS5q1K8jFtDkas3YFebu0OXrU2sGF2ayIFd/YZO00YZB+kcjXiXyTyLdEfiaQgALFanDMr48mCuRXPCMXlCPUbCIH0kl5nhOmC7RTnux2CwkliRbfbJiuiaunFZQJ255XWCtxhf0SPs9E3IfnNWEV6dLx4f4mFJ3JqaCc0SqcDdPl9IW18jhfe540v63xvBSkHtzznPAweml5zNvZcL5g4KYf69vhaLaI1uZsmK63q67Nns0jYxHPM0r52POkCYXthM8Jpfil5TGjN4Z5YX6/iMRwPFxES3g2TJfYD0AJO/PUnpclP378xvNCU0pPE14RBtBLy2NGbwwLuYhM0zEDKShxtLJnw3SJ3fiWUl+vPS9nk4sxid3kTcQ8e4BjHUO8hGfjooyt6WBgVfP54SXN6HH6vm2cvnEbp+/cxhO3boOrjJldjBtuLpOVvYSdES4S+TqRbxL5lsjPBIIXFI/u4carRIFcQQlTxAE5YVePxvOqQGyMJ/LImiKogYBlJhHdrc2GySI4XnPCzNfa80qTdvkTftIzfCCBKOR5yrykpePD/U0oOlMXVphEdBc3G6ar6ypMjLCNTO15hVx0DpCKsBtAO+Fp9Q4R3bzNycNnxYaZprCkJKJbttkwXVNXUqKZZ8crTpiL1nhecsryUc8LpM4soju1OXnkgkfGdwLWBER0tzYbpmvqakg5YVlo7XnFSDtaOl4yQhmw9bxQhK3zlpbH3NsY5noBl1JE3LKAJQYR3eDNhul6+woCoTJcT/icMiPd8xJttZaIroObKBzf583GhTBjTlxVsFz3/PCS5uxmYMzZzdCYs5uhMWc3Q486Owmbj4xsU3XuorQtXi4S+TqRbxL5lsjPBIL1Kiljzk7KRIH8bDFCF15PeMreCs2E54TtWVsiPxMIVqCkito+qRIVchUoyvLzesJTJg82E56yc3o75cP9jwXCvU8oOpMTFpikjvo8qRPl9BvzM5KcfrpYeKbNhGCkAumEJzzOXVoeM3pjWMWmNjMJh6+yjPo8mbSR1YXnSVv11xMe69YcwAgTDtsJT1mpvrQ85vPGcBl9VCjhgFdWUZsnq0RJ3YCWspNzPeHDj0KaCcEIm5m1E55iNJeWx4zdGDbjER7xcgqOoFUe9XIqTxPV8VoR1sTVEz5cZ2smBGWE03peVYQzXloe9XI2LswgP6IqrNIpapVO0at0il6lU/QqnUqs0sGNC5iKVun8zgWULSMvEvk6kW8S+ZbIzwSCnZaKVun8BgREgXyVDrG3nuCUFSITPif9fQW/B0Ks5DX7uyOwy9Hxop3fSICmifZdCPIXJzxBeWDbeJ602LWd8JI0QNfRMl0oOtMTlul0vEzndxMg6um3AEPamCcYZet/z6uS8ifkJryg7JFnecy+abspmYj8DZ+ZwrBop+NFO10kKuyKdqQ5/hOe4isazytNehjheU4qf+h4DW8MGx8SWfzLNKzh6XgNT/NEiV0Nj7L5Wj3hKfa28bxCvJ0HsE4gXqUbw6xcwEcRkRGIhmUWHa/SaZGoqJsJxAlVznrCU8xt43mlKIuqJ3x4KcvSEqizG+MFD/5Zt1P3d0qn/376q/Snq6/71bDp1ufb7c3k5WW/ufn5/wFQSwMEFAAAAAgAFU8LXYQbhIzLAAAASwEAAC0AAAB4bC9leHRlcm5hbExpbmtzL19yZWxzL2V4dGVybmFsTGluazEueG1sLnJlbHONkM1KxEAQhF8lDHjdjnvwIElO8bCgIJJ9gCbTScb5pbsXs2/vIAorePBS0FXwVVPdGwVUl5Nsrkizx5CkN5tqeQSQeaOIcsiFUk2WzBG1nrxCwdnjSnBs2wfgW4YZultmM10L/YeYl8XNNOb5EinpH2CgXYkThmeX/CvqZpoJeSXtDZyFWL4UxvyRQkYrMFGglTHeHduRxGsucI6X6K7VwKSO0TqPIThfjVSL3jEgu8MeZP9hv2Rbv3/6bjbNyfaGT/bewNDBr+2GT1BLAwQUAAAACAAVTwtdTmawVesCAAB1DwAADQAAAHhsL3N0eWxlcy54bWzdV21vmzAQ/iuIHzAIpChMBGlDijRpq6q1H/bVCSax5BdmTJf018+HKUk6X9dWzT6MqMK+y/M8d/adnRadOXB6u6PUBHvBZbcMd8a0H6Oo2+yoIN0H1VJpPY3Sghg71duoazUldQcgwaMkjrNIECbDspC9WAnTBRvVS7MMZ5MpcK8vtTVm8zBwdJWq6TKMP8RhVBbRCC6LRskjRxY6g2Uiggb3hC/DinC21gxQDRGMH5w5AcNGcaUDY4OnEIC1dA/OPXMzyGvkEUwqPWg7hdforEfS99a8Y4J2wTX9FXxXgsiBfkd0ZzfIicfzp/H8LYb5BeQuSu5y0dv1MlzZJ7bPxTVfuJ+n0sML6pVxPtVrEjpDWbTEGKrlyk4GzGD8wxWM47tDa8W2mhxmyVV4BAwvK7JWuqb6rLWcqSw4bYwFaLbdwduoFvJRxihhBzUjWyXJEMMjYhxY2g3l/BbOgB/NGfe+OWnZGBpWTkMb0Dh0NG4C/KdsjvuENn8TbdCye2U+9zYbOcx/9srQG00bth/m+2bSx9gTnJ20LT984mwrBXW5v1iwLMgjLtgpzR6sGtTIxhqoa/F986ag3ppy8q/Y0/dnH6+Fi0V/svLzi67N1UVjn70LezS25kn/n3X/ZA3gYF2G13Bd8yNFsO4ZN0x6Ot9yGrK2vynOSC2kpg3pubmbnPasnsbfaM16kUzfuoFcxm8dx1/hhJtl0xFttZis6Z7W1Ti1d0Z1vDzi8QHAU89qePweDON8fg/4MB0sAgzjUJjO/5TPAs3H+bDYFl7PAsUsUIxD+TzV8MF0/JjcPv5M8zxNswxb0aryRlBh65Zl8Odnw2IDBKYDSq9ba3y38Qp5vg6wPX2uQrBM8UrEMsXXGjz+dQNEnvt3G9MBBLYLWO2Avl8HasqPSVPYVSw2rINxT55jHqhFf41mGbI6GXz8+4N1SZrmud8DPn8EaYp5oBtxDxYBxIB50nS4B5/cR9HjPRUd/9EufwNQSwMEFAAAAAgAFU8LXZeKuxzAAAAAEwIAAAsAAABfcmVscy8ucmVsc52SuW7DMAxAf8XQnjAH0CGIM2XxFgT5AVaiD9gSBYpFnb+v2qVxkAsZeT08EtweaUDtOKS2i6kY/RBSaVrVuAFItiWPac6RQq7ULB41h9JARNtjQ7BaLD5ALhlmt71kFqdzpFeIXNedpT3bL09Bb4CvOkxxQmlISzMO8M3SfzL38ww1ReVKI5VbGnjT5f524EnRoSJYFppFydOiHaV/Hcf2kNPpr2MitHpb6PlxaFQKjtxjJYxxYrT+NYLJD+x+AFBLAwQUAAAACAAVTwtd4lm4aGoBAADsAgAADwAAAHhsL3dvcmtib29rLnhtbK2SSU/EMAyF/0qVO7RTEMtoOhcQi4QAAYJzpnGnFlkqx0MHfj1uqrIIDhw4Oe+5cj6/dNEHel6F8JxtnfWxUi1zN8/zWLfgdNwNHXjpNIGcZpG0zmNHoE1sAdjZvCyKg9xp9Gq5mGbdUr5cDIdHhD5++oPMXjDiCi3ya6XS2YLKHHp0+AamUoXKYhv6i0D4Fjxre19TsLZSs7HxCMRY/7DvB54HvYrJ2T6hN6Gv1M6slIGv32Wf1BMabitVFod7H94F4LplGTEr9sVkvbrTjKFSB4XIBilyuihh6prxBeTOUW04nKFloFPNcE5h06FfDzQSRv4ljZTcVMfY5/SX4EPTYA2nod448DwmT2AHQB9b7KLKvHZQqYSYcpF6acaMWKi+JE5zlAZdmpFvgoKtLOC1vYMGCHwNv3n/DP2BUiaU3xAMNOjBXMt6Ub6R569vKRtKWu9oVpTH8j4ba0/Eu/FXQZsp+unvW74DUEsDBBQAAAAIABVPC12dRm95wQAAAJYCAAAaAAAAeGwvX3JlbHMvd29ya2Jvb2sueG1sLnJlbHO1kj0OgzAMha+CcoAaaNWhAqYuSJ0qLhCB+REhiWJXhduXwgCROnRhip7tvPfJcvJEJbkzmtrOUjAOSlMqWmZ7A6CyxUHSyVjUc6c2bpA8S9eAlWUvG4Q4DK/g9h4iS/aeQTFZ/MfR1HVX4t2UrwE1/zCGt3E9tYgsgkK6BjkVMKqtTLA80Wl2FkFepcLlVSTgMCAcGZ2W6tHp3mfad/w5ny4+kI54Ukgb16q9+POB8Tz/xS19kWvRX8HlywDeGWYfUEsDBBQAAAAIABVPC11Fu7/ySwEAANQFAAATAAAAW0NvbnRlbnRfVHlwZXNdLnhtbL2UzW7CMAzHX6XqFbVhHHaYgMu227Rx2AtkiUuj5ktx+Hr7OS2gMbECKtqlrhP7/7PjKNPPnQfMtkZbnOV1jP6JMRQ1GI6l82Bpp3LB8EhuWDLPRcOXwCbj8SMTzkawsYhJI59PX6DiKx2z1y0to3J2lgfQmGfPXWBizXLuvVaCR9pnayt/UYo9oaTMNgZr5XFEATk7S0g7fwP2eR9rCEFJyBY8xHduKIptNcO404Blv8SZGl1VKQHSiZWhlBJ9AC6xBohGl53oqJ8c6YSh+z4M5rcyfUCKXATnkSYW4HbcYSQpu/AkBCGq/haPRJIe3B+kaUuQV7LpeDcuNO08kLVm+Bmfzviof6EOGfhG2SUefobXsRe6wBU1udiZuzGp71bwBvbkn9nppgTL9ZuyDZ54974BP7WvuIxfzjX3fmeSLQ1X9sBn7WM+/wZQSwECFAMUAAAACAAVTwtdRsdNSJUAAADNAAAAEAAAAAAAAAAAAAAAgAEAAAAAZG9jUHJvcHMvYXBwLnhtbFBLAQIUAxQAAAAIABVPC10muw0XEAEAAHACAAARAAAAAAAAAAAAAACAAcMAAABkb2NQcm9wcy9jb3JlLnhtbFBLAQIUAxQAAAAIABVPC13BFxC+kgYAAMYgAAATAAAAAAAAAAAAAACAAQICAAB4bC90aGVtZS90aGVtZTEueG1sUEsBAhQDFAAAAAgAFU8LXb/9bw5XCwAAY2IAABgAAAAAAAAAAAAAAICBxQgAAHhsL3dvcmtzaGVldHMvc2hlZXQxLnhtbFBLAQIUAxQAAAAIABVPC13N4o62ZQEAAJ0FAAAYAAAAAAAAAAAAAACAAVIUAAB4bC9kcmF3aW5ncy9kcmF3aW5nMS54bWxQSwECFAMUAAAACAAVTwtdMLsGyJQAAAB2AQAAIwAAAAAAAAAAAAAAgAHtFQAAeGwvZHJhd2luZ3MvX3JlbHMvZHJhd2luZzEueG1sLnJlbHNQSwECFAMUAAAACAAVTwtdd/cihY0AAADzAAAAIwAAAAAAAAAAAAAAgAHCFgAAeGwvd29ya3NoZWV0cy9fcmVscy9zaGVldDEueG1sLnJlbHNQSwECFAMUAAAACAAVTwtd7SyAGLAFAAD0MgAAFAAAAAAAAAAAAAAAgAGQFwAAeGwvY2hhcnRzL2NoYXJ0MS54bWxQSwECFAMUAAAACAAVTwtdrsAaN50EAAAvHgAAFAAAAAAAAAAAAAAAgAFyHQAAeGwvY2hhcnRzL2NoYXJ0Mi54bWxQSwECFAMUAAAACAAVTwtdvggw9YkUAACMjQAAIgAAAAAAAAAAAAAAgAFBIgAAeGwvZXh0ZXJuYWxMaW5rcy9leHRlcm5hbExpbmsxLnhtbFBLAQIUAxQAAAAIABVPC12EG4SMywAAAEsBAAAtAAAAAAAAAAAAAACAAQo3AAB4bC9leHRlcm5hbExpbmtzL19yZWxzL2V4dGVybmFsTGluazEueG1sLnJlbHNQSwECFAMUAAAACAAVTwtdTmawVesCAAB1DwAADQAAAAAAAAAAAAAAgAEgOAAAeGwvc3R5bGVzLnhtbFBLAQIUAxQAAAAIABVPC12XirscwAAAABMCAAALAAAAAAAAAAAAAACAATY7AABfcmVscy8ucmVsc1BLAQIUAxQAAAAIABVPC13iWbhoagEAAOwCAAAPAAAAAAAAAAAAAACAAR88AAB4bC93b3JrYm9vay54bWxQSwECFAMUAAAACAAVTwtdnUZvecEAAACWAgAAGgAAAAAAAAAAAAAAgAG2PQAAeGwvX3JlbHMvd29ya2Jvb2sueG1sLnJlbHNQSwECFAMUAAAACAAVTwtdRbu/8ksBAADUBQAAEwAAAAAAAAAAAAAAgAGvPgAAW0NvbnRlbnRfVHlwZXNdLnhtbFBLBQYAAAAAEAAQAFUEAAArQAAAAAA="

def ensure_template_exists():
    if not os.path.isfile(TEMPLATE_PATH):
        try:
            with open(TEMPLATE_PATH, "wb") as f:
                f.write(base64.b64decode(DPPH_TEMPLATE_B64))
        except Exception as e:
            logging.warning(f"Shablonni yaratishda xatolik: {e}")

ensure_template_exists()

# =====================================================================
# MATEMATIK HISOB-KITOBLAR (DPPH / IC50)
# =====================================================================

TIMES = [5, 10, 15, 20, 25, 30]
VOLUMES = [25, 50, 75, 100]

SAMPLE_EXAMPLE_TEXT = """1.048
0.816
0.601
0.539
0.454
0.675
0.302
0.236
0.261
0.614
0.253
0.225
0.263
0.584
0.237
0.219
0.260
0.558
0.225
0.215
0.258
0.540
0.218
0.213
0.259"""

def parse_absorbance_values(raw_text: str) -> List[float]:
    if not raw_text:
        return []
    raw = raw_text.strip().replace(",", ".")
    for sep in ["\t", ";", "\n", "\r", " "]:
        raw = raw.replace(sep, " ")
    tokens = raw.split()
    values = []
    for token in tokens:
        try:
            values.append(float(token))
        except ValueError:
            pass
    return values

def calculate_dpph(
    values: List[float],
    control: Optional[float] = None,
    selected_time: int = 30
) -> Dict[str, Any]:
    if len(values) == 25:
        control = values[0]
        sample_values = values[1:]
    elif len(values) == 24:
        if control is None or control <= 0:
            raise ValueError("Control qiymati kiritilmagan yoki noto'g'ri!")
        sample_values = values
    else:
        raise ValueError(
            f"{len(values)} ta qiymat topildi. 24 ta sample Abs yoki 25 ta (1 control + 24 sample) kerak."
        )

    if control <= 0:
        raise ValueError("DPPH Control Abs qiymati 0 dan katta bo'lishi kerak!")

    rows = []
    for time_index, minute in enumerate(TIMES):
        for volume_index, volume in enumerate(VOLUMES):
            index = time_index * len(VOLUMES) + volume_index
            absorbance = sample_values[index]
            activity = ((control - absorbance) / control) * 100.0
            rows.append({
                "Minute": minute,
                "DPPH": control,
                "Volume": volume,
                "Abs": absorbance,
                "Antiradical Activity (%)": activity
            })

    selected_rows = [r for r in rows if r["Minute"] == selected_time]
    if not selected_rows:
        selected_time = 30
        selected_rows = [r for r in rows if r["Minute"] == selected_time]

    y_activities = [r["Antiradical Activity (%)"] for r in selected_rows]
    regression_x = np.array([0.0, 25.0, 50.0, 75.0, 100.0], dtype=float)
    regression_y = np.array([0.0] + y_activities, dtype=float)

    slope, intercept = np.polyfit(regression_x, regression_y, 1)
    predicted = slope * regression_x + intercept
    ss_res = np.sum((regression_y - predicted) ** 2)
    ss_tot = np.sum((regression_y - np.mean(regression_y)) ** 2)
    r2 = 1 - ss_res / ss_tot if ss_tot != 0 else 0.0

    if slope != 0:
        ic50 = (50.0 - intercept) / slope
    else:
        ic50 = None

    return {
        "control": control,
        "sample_values": sample_values,
        "selected_time": selected_time,
        "rows": rows,
        "selected_rows": selected_rows,
        "slope": float(slope),
        "intercept": float(intercept),
        "r2": float(r2),
        "ic50": float(ic50) if ic50 is not None else None
    }

# =====================================================================
# GRAFIKLAR (MATPLOTLIB)
# =====================================================================

def generate_activity_vs_time_chart(rows: List[Dict[str, Any]], plant_name: str = "Namuna") -> io.BytesIO:
    volumes = [25, 50, 75, 100]
    plt.figure(figsize=(9, 5.5), dpi=150)
    colors = ["#2563EB", "#16A34A", "#EA580C", "#DC2626"]
    markers = ["o", "s", "^", "D"]

    for idx, volume in enumerate(volumes):
        data = [r for r in rows if r["Volume"] == volume]
        x = [r["Minute"] for r in data]
        y = [r["Antiradical Activity (%)"] for r in data]
        plt.plot(
            x, y,
            marker=markers[idx % len(markers)],
            color=colors[idx % len(colors)],
            linewidth=2,
            markersize=6,
            label=f"{volume} µL"
        )

    plt.xlabel("Time (min)", fontsize=11, fontweight="bold")
    plt.ylabel("Antiradical activity (%)", fontsize=11, fontweight="bold")
    plt.title(f"DPPH Antiradical Activity vs Time\n({plant_name})", fontsize=12, fontweight="bold", pad=12)
    plt.grid(True, linestyle="--", alpha=0.5)
    plt.legend(frameon=True, facecolor="white", edgecolor="none", shadow=True)
    plt.tight_layout()

    buf = io.BytesIO()
    plt.savefig(buf, format="png", dpi=150)
    plt.close()
    buf.seek(0)
    return buf

def generate_ic50_regression_chart(
    selected_rows: List[Dict[str, Any]],
    slope: float,
    intercept: float,
    ic50: Optional[float],
    selected_time: int = 30,
    plant_name: str = "Namuna"
) -> io.BytesIO:
    x = np.array([0.0] + [r["Volume"] for r in selected_rows], dtype=float)
    y = np.array([0.0] + [r["Antiradical Activity (%)"] for r in selected_rows], dtype=float)

    max_x = max(max(x), (ic50 * 1.1) if (ic50 and ic50 > 0 and ic50 < 300) else 100.0)
    x_line = np.linspace(0.0, max_x, 150)
    y_line = slope * x_line + intercept

    plt.figure(figsize=(9, 5.5), dpi=150)
    plt.scatter(x, y, color="#1E40AF", s=80, zorder=5, label="Experimental data")
    plt.plot(x_line, y_line, color="#2563EB", linewidth=2, label=f"Linear regression (y = {slope:.4f}x + {intercept:.4f})")
    plt.axhline(50, color="#DC2626", linestyle="--", linewidth=1.5, label="50% inhibition")

    if ic50 is not None and ic50 > 0:
        plt.axvline(ic50, color="#16A34A", linestyle="--", linewidth=1.5, label=f"IC₅₀ = {ic50:.2f} µL")
        plt.scatter([ic50], [50], color="#DC2626", s=100, zorder=6)

    plt.xlabel("Volume (µL)", fontsize=11, fontweight="bold")
    plt.ylabel("Antiradical activity (%)", fontsize=11, fontweight="bold")
    plt.title(f"IC₅₀ Determination at {selected_time} min\n({plant_name})", fontsize=12, fontweight="bold", pad=12)
    plt.grid(True, linestyle="--", alpha=0.5)
    plt.legend(frameon=True, facecolor="white", edgecolor="none", shadow=True)
    plt.tight_layout()

    buf = io.BytesIO()
    plt.savefig(buf, format="png", dpi=150)
    plt.close()
    buf.seek(0)
    return buf

# =====================================================================
# EXCEL EKSPORT (DINAMIK TAHRIRLANADIGAN SHABLON)
# =====================================================================

def export_dpph_excel(
    plant_name: str,
    calc_result: Dict[str, Any],
    template_path: str,
    output_path: str
) -> str:
    ensure_template_exists()
    if not os.path.isfile(template_path):
        raise FileNotFoundError(f"DPPH shablon fayli topilmadi: {template_path}")

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    shutil.copy2(template_path, output_path)

    wb = load_workbook(output_path)
    ws = wb.active

    if not plant_name:
        plant_name = "Noma'lum o'simlik"

    ws["A1"] = plant_name
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["A1"].font = Font(name="Calibri", bold=True, size=13)
    ws.row_dimensions[1].height = 32

    try:
        if hasattr(ws, "_charts"):
            if len(ws._charts) >= 1:
                ws._charts[0].title = plant_name
            if len(ws._charts) >= 2:
                ws._charts[1].title = plant_name
    except Exception:
        pass

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass

    control = float(calc_result["control"])
    rows = calc_result["rows"]
    selected_time = int(calc_result.get("selected_time", 30))

    times = [0, 5, 10, 15, 20, 25, 30]
    volumes = [25, 50, 75, 100]

    font_bold = Font(name="Calibri", bold=True, size=11)
    font_regular = Font(name="Calibri", bold=False, size=11)
    font_red_bold = Font(name="Calibri", bold=True, size=11, color="FF0000")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF")
    )

    # Sarlavhalar
    headers = {"A2": "Daqiqa", "B2": "DPPH", "C2": "25 mkl", "D2": None, "E2": "50 mkl", "F2": None, "G2": "75 mkl", "H2": None, "I2": "100 mkl", "J2": None}
    for coord, text in headers.items():
        ws[coord] = text
        ws[coord].font = font_bold
        ws[coord].alignment = align_center
        ws[coord].border = thin_border

    # 0-30 daqiqalar
    for excel_row, minute in enumerate(times, start=3):
        c_min = ws.cell(excel_row, 1, minute)
        c_min.font = font_bold if minute == selected_time else font_regular
        c_min.alignment = align_center
        c_min.border = thin_border

        if excel_row == 3:
            c_ctrl = ws.cell(excel_row, 2, control)
        else:
            c_ctrl = ws.cell(excel_row, 2, "=$B$3")
        c_ctrl.font = font_regular
        c_ctrl.alignment = align_center
        c_ctrl.number_format = "0.000"
        c_ctrl.border = thin_border

    # 0 daqiqa (Control)
    for col_idx, form_col in [(3, 4), (5, 6), (7, 8), (9, 10)]:
        c_abs = ws.cell(3, col_idx, "=$B$3")
        c_abs.font = font_regular
        c_abs.alignment = align_center
        c_abs.number_format = "0.000"
        c_abs.border = thin_border

        abs_letter = {3: "C", 5: "E", 7: "G", 9: "I"}[col_idx]
        c_aa = ws.cell(3, form_col, f"=(B3-{abs_letter}3)/B3*100")
        c_aa.font = font_bold
        c_aa.alignment = align_right
        c_aa.number_format = "0.00"
        c_aa.border = thin_border

    # 5-30 daqiqalar
    column_map = {25: (3, 4), 50: (5, 6), 75: (7, 8), 100: (9, 10)}
    for excel_row, minute in zip(range(4, 10), [5, 10, 15, 20, 25, 30]):
        for volume in volumes:
            matching = [r for r in rows if r["Volume"] == volume and r["Minute"] == minute]
            if not matching:
                continue
            data = matching[0]
            abs_col, aa_col = column_map[volume]

            c_abs = ws.cell(excel_row, abs_col, data["Abs"])
            c_abs.font = font_regular
            c_abs.alignment = align_center
            c_abs.number_format = "0.000" if len(str(data["Abs"]).split(".")[-1]) > 2 else "0.00"
            c_abs.border = thin_border

            abs_letter = {3: "C", 5: "E", 7: "G", 9: "I"}[abs_col]
            c_aa = ws.cell(excel_row, aa_col, f"=(B{excel_row}-{abs_letter}{excel_row})/B{excel_row}*100")
            c_aa.font = font_bold
            c_aa.alignment = align_right
            c_aa.number_format = "0.00"
            c_aa.border = thin_border

    # K va L ustunlari
    target_row = 3 + (selected_time // 5)
    if target_row < 4 or target_row > 9:
        target_row = 9

    ws["K2"] = None
    ws["L2"] = None

    ws["K5"] = 0
    ws["K5"].font = font_regular
    ws["K5"].alignment = align_center
    ws["K5"].border = thin_border

    ws["L5"] = 0.0
    ws["L5"].font = font_bold
    ws["L5"].alignment = align_right
    ws["L5"].number_format = "0.00"
    ws["L5"].border = thin_border

    regression_links = [
        (6, 25, f"=D{target_row}"),
        (7, 50, f"=F{target_row}"),
        (8, 75, f"=H{target_row}"),
        (9, 100, f"=J{target_row}")
    ]
    for excel_row, volume, formula in regression_links:
        c_k = ws.cell(excel_row, 11, volume)
        c_k.font = font_regular
        c_k.alignment = align_center
        c_k.border = thin_border

        c_l = ws.cell(excel_row, 12, formula)
        c_l.font = font_bold
        c_l.alignment = align_right
        c_l.number_format = "0.00"
        c_l.border = thin_border

    # N va O ustunlari
    ws["N5"] = "m"
    ws["N5"].font = font_bold
    ws["N5"].alignment = align_center
    ws["N5"].border = thin_border
    ws["O5"] = "=SLOPE(L5:L9,K5:K9)"
    ws["O5"].font = font_regular
    ws["O5"].alignment = align_right
    ws["O5"].number_format = "0.0000"
    ws["O5"].border = thin_border

    ws["N6"] = "b"
    ws["N6"].font = font_bold
    ws["N6"].alignment = align_center
    ws["N6"].border = thin_border
    ws["O6"] = "=INTERCEPT(L5:L9,K5:K9)"
    ws["O6"].font = font_regular
    ws["O6"].alignment = align_right
    ws["O6"].number_format = "0.0000"
    ws["O6"].border = thin_border

    ws["N7"] = "y"
    ws["N7"].font = font_bold
    ws["N7"].alignment = align_center
    ws["N7"].border = thin_border
    ws["O7"] = 50
    ws["O7"].font = font_bold
    ws["O7"].alignment = align_right
    ws["O7"].border = thin_border

    ws["N8"] = "x=(y-b)/m"
    ws["N8"].font = font_bold
    ws["N8"].alignment = align_center
    ws["N8"].border = thin_border
    ws["O8"] = "=(O7-O6)/O5"
    ws["O8"].font = font_red_bold
    ws["O8"].alignment = align_right
    ws["O8"].number_format = "0.00"
    ws["O8"].border = thin_border

    ws["N10"] = "IC₅₀ time"
    ws["N10"].font = font_bold
    ws["N10"].alignment = align_center
    ws["N10"].border = thin_border
    ws["O10"] = selected_time
    ws["O10"].font = font_bold
    ws["O10"].alignment = align_center
    ws["O10"].border = thin_border

    col_widths = {"A": 8, "B": 9, "C": 8, "D": 8, "E": 8, "F": 8, "G": 8, "H": 8, "I": 8, "J": 8, "K": 8, "L": 9, "M": 3, "N": 12, "O": 10}
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(output_path)
    return output_path

# =====================================================================
# TELEGRAM BOT KLAVIATURALARI VA HOLATLARI (FSM)
# =====================================================================

class DPPHStates(StatesGroup):
    waiting_plant_name = State()
    waiting_time = State()
    waiting_abs_values = State()

def get_main_keyboard() -> ReplyKeyboardMarkup:
    kb = [
        [KeyboardButton(text="🧬 Yangi DPPH tahlili")],
        [KeyboardButton(text="📋 Namuna bilan sinash (Demo)")],
        [KeyboardButton(text="ℹ️ Yo'riqnoma va Ma'lumot")]
    ]
    return ReplyKeyboardMarkup(keyboard=kb, resize_keyboard=True)

def get_cancel_keyboard() -> ReplyKeyboardMarkup:
    kb = [[KeyboardButton(text="❌ Bekor qilish")]]
    return ReplyKeyboardMarkup(keyboard=kb, resize_keyboard=True)

def get_skip_or_cancel_keyboard() -> ReplyKeyboardMarkup:
    kb = [
        [KeyboardButton(text="⏭ O'tkazib yuborish")],
        [KeyboardButton(text="❌ Bekor qilish")]
    ]
    return ReplyKeyboardMarkup(keyboard=kb, resize_keyboard=True)

def get_time_selection_inline_keyboard(current_time: int = 30) -> InlineKeyboardMarkup:
    times = [5, 10, 15, 20, 25, 30]
    buttons = []
    row = []
    for t in times:
        prefix = "✅ " if t == current_time else ""
        row.append(InlineKeyboardButton(text=f"{prefix}{t} min", callback_data=f"set_time:{t}"))
        if len(row) == 3:
            buttons.append(row)
            row = []
    if row:
        buttons.append(row)
    buttons.append([InlineKeyboardButton(text="Davom etish ➡️", callback_data="confirm_time")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)

def get_example_fill_inline_keyboard() -> InlineKeyboardMarkup:
    buttons = [[InlineKeyboardButton(text="📋 Namunaviy qiymatlar bilan to'ldirish", callback_data="fill_example_abs")]]
    return InlineKeyboardMarkup(inline_keyboard=buttons)

# =====================================================================
# TELEGRAM BOT HANDLERS
# =====================================================================

router = Router()

@router.message(CommandStart())
async def cmd_start(message: Message, state: FSMContext):
    await state.clear()
    welcome_text = (
        "👋 **Assalomu alaykum!**\n\n"
        "🌱 **Plant Chemical Analysis — DPPH / IC₅₀ Tahlil Botiga xush kelibsiz!**\n\n"
        "Ushbu bot yordamida spektrofotometr ma'lumotlari asosida "
        "o'simliklarning **antiradikal faolligi** va **IC₅₀** ko'rsatkichini "
        "avtomatik hisoblab, grafiklar va tahrirlanadigan Excel hisobotini olishingiz mumkin.\n\n"
        "Quyidagi menyudan kerakli bo'limni tanlang 👇"
    )
    await message.answer(welcome_text, reply_markup=get_main_keyboard(), parse_mode="Markdown")

@router.message(F.text == "❌ Bekor qilish")
@router.message(Command("cancel"))
async def cmd_cancel(message: Message, state: FSMContext):
    await state.clear()
    await message.answer("❌ Amaliyot bekor qilindi. Bosh menyudasiz.", reply_markup=get_main_keyboard())

@router.message(F.text == "ℹ️ Yo'riqnoma va Ma'lumot")
async def cmd_info(message: Message):
    info_text = (
        "📖 **DPPH / IC₅₀ Botidan foydalanish bo'yicha yo'riqnoma:**\n\n"
        "1️⃣ **Namuna nomi:** O'simlik yoki ekstrakt nomini kiriting (masalan: *Curcuma longa L.*).\n"
        "2️⃣ **IC₅₀ vaqti:** Regressiya hisoblanadigan vaqtni tanlang (standart: **30 min**).\n"
        "3️⃣ **Spektrofotometr ABS qiymatlari:**\n"
        "   • Hammasi bo'lib **25 ta** (1 ta Control + 24 ta Sample) yoki **24 ta** qiymat kiritiladi.\n"
        "   • Tartib: 5, 10, 15, 20, 25, 30 daqiqa uchun har 4 tadan (25, 50, 75, 100 µL).\n"
        "   • Qiymatlarni shunchaki bo'shliq yoki yangi qator bilan nusxalab yuborish kifoya.\n\n"
        "📊 **Natijada nima olasiz?**\n"
        "   ✅ Chiziqli regressiya formulasi (y = mx + b) va R²\n"
        "   ✅ Hisoblangan IC₅₀ (µL)\n"
        "   ✅ 2 ta sifatli grafik (Faollik dinamikasi va IC₅₀ regressiyasi)\n"
        "   ✅ To'liq tahrirlanadigan Excel (.xlsx) hisoboti"
    )
    await message.answer(info_text, reply_markup=get_main_keyboard(), parse_mode="Markdown")

@router.message(F.text == "🧬 Yangi DPPH tahlili")
async def start_dpph_analysis(message: Message, state: FSMContext):
    await state.set_state(DPPHStates.waiting_plant_name)
    await message.answer(
        "🌿 **1-qadam:** O'simlik / namuna nomini kiriting:\n\n"
        "*(Masalan: Curcuma longa L. (kurkuma))*\n\n"
        "Agar nom kiritishni xohlamasangiz, pastdagi **'⏭ O'tkazib yuborish'** tugmasini bosing.",
        reply_markup=get_skip_or_cancel_keyboard(),
        parse_mode="Markdown"
    )

@router.message(DPPHStates.waiting_plant_name)
async def process_plant_name(message: Message, state: FSMContext):
    if message.text == "⏭ O'tkazib yuborish":
        plant_name = "Noma'lum o'simlik"
    else:
        plant_name = message.text.strip()
    await state.update_data(plant_name=plant_name, selected_time=30)
    await state.set_state(DPPHStates.waiting_time)
    await message.answer(
        f"✅ O'simlik nomi: **{plant_name}**\n\n"
        "⏱ **2-qadam:** IC₅₀ hisoblash uchun vaqtni tanlang (standart: **30 min**):",
        reply_markup=get_time_selection_inline_keyboard(30),
        parse_mode="Markdown"
    )

@router.callback_query(DPPHStates.waiting_time, F.data.startswith("set_time:"))
async def callback_set_time(callback: CallbackQuery, state: FSMContext):
    selected = int(callback.data.split(":")[1])
    await state.update_data(selected_time=selected)
    await callback.message.edit_reply_markup(reply_markup=get_time_selection_inline_keyboard(selected))
    await callback.answer(f"Tanlandi: {selected} min")

@router.callback_query(DPPHStates.waiting_time, F.data == "confirm_time")
async def callback_confirm_time(callback: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    selected_time = data.get("selected_time", 30)
    await state.set_state(DPPHStates.waiting_abs_values)
    await callback.message.delete()
    prompt_text = (
        f"⏱ Tanlangan vaqt: **{selected_time} min**\n\n"
        "🧪 **3-qadam: Spektrofotometr ABS qiymatlarini yuboring.**\n\n"
        "Spektrofotometrdan **25 ta qiymat** (1-chi Control + 24 ta Sample) "
        "yoki 24 ta Sample qiymatlarini nusxalab, shu yerga xabar qilib yuboring.\n\n"
        "*(Tartib: 5, 10, 15, 20, 25, 30 minutlar uchun har 4 tadan: "
        "25 µL, 50 µL, 75 µL, 100 µL)*\n\n"
        "Agar test qilib ko'rmoqchi bo'lsangiz, pastdagi tugmani bosing 👇"
    )
    await callback.message.answer(prompt_text, reply_markup=get_example_fill_inline_keyboard(), parse_mode="Markdown")

@router.callback_query(DPPHStates.waiting_abs_values, F.data == "fill_example_abs")
async def callback_fill_example(callback: CallbackQuery, state: FSMContext):
    await callback.answer("Namunaviy qiymatlar yuklanmoqda...")
    await execute_calculation(callback.message, state, SAMPLE_EXAMPLE_TEXT)

@router.message(DPPHStates.waiting_abs_values, F.text)
async def process_abs_input(message: Message, state: FSMContext):
    await execute_calculation(message, state, message.text)

@router.message(F.text == "📋 Namuna bilan sinash (Demo)")
async def run_demo_analysis(message: Message, state: FSMContext):
    await state.clear()
    await state.update_data(plant_name="Curcuma longa L. (Demo)", selected_time=30)
    await message.answer(
        "🚀 **Demo rejim:** Namunaviy ma'lumotlar bilan hisob-kitob bajarilmoqda...",
        reply_markup=get_main_keyboard(),
        parse_mode="Markdown"
    )
    await execute_calculation(message, state, SAMPLE_EXAMPLE_TEXT)

async def execute_calculation(message: Message, state: FSMContext, raw_text: str):
    data = await state.get_data()
    plant_name = data.get("plant_name", "Curcuma longa L.")
    selected_time = data.get("selected_time", 30)

    values = parse_absorbance_values(raw_text)
    if not values:
        await message.answer("⚠️ Hech qanday raqam aniqlanmadi. Iltimos, spektrofotometr Abs qiymatlarini to'g'ri kiriting.", reply_markup=get_cancel_keyboard())
        return

    if len(values) not in (24, 25):
        await message.answer(
            f"⚠️ **{len(values)} ta** qiymat topildi.\n\n"
            f"Tahlil uchun **25 ta** (1-chi Control + 24 ta Sample) yoki **24 ta** qiymat bo'lishi kerak.\n"
            f"Iltimos, qaytadan tekshirib nusxa ko'chiring.",
            reply_markup=get_cancel_keyboard(),
            parse_mode="Markdown"
        )
        return

    status_msg = await message.answer("⏳ **Hisoblash va grafiklar tayyorlanmoqda...**", parse_mode="Markdown")

    try:
        calc_result = calculate_dpph(values, selected_time=selected_time)
        control = calc_result["control"]
        slope = calc_result["slope"]
        intercept = calc_result["intercept"]
        r2 = calc_result["r2"]
        ic50 = calc_result["ic50"]
        ic50_str = f"{ic50:.4f} µL" if ic50 is not None else "Hisoblanmadi"

        report_text = (
            f"✅ **DPPH / IC₅₀ Tahlil Natijalari**\n\n"
            f"🌿 **O'simlik / Namuna:** `{plant_name}`\n"
            f"🧪 **DPPH Control Abs:** `{control:.4f}`\n"
            f"⏱ **IC₅₀ hisoblangan vaqt:** `{selected_time} min`\n\n"
            f"📊 **Chiziqli Regressiya (Linear Regression):**\n"
            f"• Tenglama: `y = {slope:.6f}x + {intercept:.6f}`\n"
            f"• Qiyalik (m / slope): `{slope:.6f}`\n"
            f"• Siljish (b / intercept): `{intercept:.6f}`\n"
            f"• Determinatsiya koeffitsiyenti (R²): `{r2:.6f}`\n\n"
            f"🎯 **IC₅₀ (50% ingibitsiya konsentratsiyasi):**\n"
            f"👉 **{ic50_str}**\n\n"
            f"Grafiklar va tahrirlanadigan Excel hisoboti quyida yuklanmoqda 👇"
        )
        await status_msg.edit_text(report_text, parse_mode="Markdown")

        chart1_buf = generate_activity_vs_time_chart(calc_result["rows"], plant_name)
        photo1 = BufferedInputFile(chart1_buf.read(), filename=f"dpph_activity_{selected_time}min.png")
        await message.answer_photo(photo1, caption=f"📈 **Antiradikal faollik dinamikasi (Vaqt bo'yicha)**\nNamuna: {plant_name}", parse_mode="Markdown")

        chart2_buf = generate_ic50_regression_chart(calc_result["selected_rows"], slope, intercept, ic50, selected_time, plant_name)
        photo2 = BufferedInputFile(chart2_buf.read(), filename=f"ic50_regression_{selected_time}min.png")
        await message.answer_photo(photo2, caption=f"📉 **IC₅₀ Chiziqli Regressiya Grafigi**\nIC₅₀ = {ic50_str} (R² = {r2:.4f})", parse_mode="Markdown")

        safe_name = "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in plant_name).strip() or "Sample"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        excel_filename = f"DPPH_{safe_name}_{timestamp}.xlsx"
        excel_path = os.path.join(OUTPUT_DIR, excel_filename)

        export_dpph_excel(
            plant_name=plant_name,
            calc_result=calc_result,
            template_path=TEMPLATE_PATH,
            output_path=excel_path
        )

        excel_doc = FSInputFile(excel_path, filename=excel_filename)
        await message.answer_document(
            excel_doc,
            caption="📑 **Tahrirlanadigan Excel hisoboti (.xlsx)**\nBarcha formulalar va hisob-kitoblar kiritilgan.",
            reply_markup=get_main_keyboard()
        )
        await state.clear()

    except Exception as e:
        await message.answer(f"❌ Hisoblash jarayonida xatolik yuz berdi: `{str(e)}`", reply_markup=get_main_keyboard(), parse_mode="Markdown")
        await state.clear()

# =====================================================================
# RENDER UCHUN WEB PING SERVERI VA ISHGA TUSHIRISH
# =====================================================================

async def start_web_server():
    port_str = os.getenv("PORT")
    if port_str:
        try:
            port = int(port_str)
            app = web.Application()
            async def handle_ping(request):
                return web.Response(text="🤖 Plant Chemical Analysis Bot is running 24/7!")
            app.router.add_get("/", handle_ping)
            app.router.add_get("/health", handle_ping)
            runner = web.AppRunner(app)
            await runner.setup()
            site = web.TCPSite(runner, "0.0.0.0", port)
            await site.start()
            logger.info(f"🌐 Cloud Web Server {port}-portda ishga tushdi va port ochiq.")
        except Exception as e:
            logger.warning(f"Web serverni ishga tushirishda xatolik: {e}")

async def main():
    await start_web_server()

    if not BOT_TOKEN or BOT_TOKEN == "YOUR_BOT_TOKEN_HERE":
        logger.error(
            "\n" + "=" * 60 + "\n"
            "[XATOLIK] BOT_TOKEN kiritilmagan!\n"
            "Iltimos, Render Environment Variables bo'limiga BOT_TOKEN qo'shing:\n"
            "BOT_TOKEN=8939401003:AAGO2fMFsR489ZisMQO4DGZBETAt6W478ag\n"
            + "=" * 60
        )
        while True:
            await asyncio.sleep(3600)
        return

    logger.info("🤖 Plant Chemical Analysis (DPPH / IC50) bot ishga tushmoqda...")

    bot = Bot(token=BOT_TOKEN, default=DefaultBotProperties(parse_mode=ParseMode.MARKDOWN))
    dp = Dispatcher(storage=MemoryStorage())
    dp.include_router(router)

    try:
        await bot.delete_webhook(drop_pending_updates=True)
        bot_info = await bot.get_me()
        logger.info(f"✅ Bot muvaffaqiyatli ulandi: @{bot_info.username} ({bot_info.first_name})")
    except Exception as e:
        logger.error(f"Telegram API ga ulanishda xatolik: {e}")

    try:
        await dp.start_polling(bot)
    finally:
        await bot.session.close()
        logger.info("Bot to'xtatildi.")

if __name__ == "__main__":
    try:
        asyncio.run(main())
    except (KeyboardInterrupt, SystemExit):
        logger.info("Bot foydalanuvchi tomonidan to'xtatildi.")
