import os
import unittest
import openpyxl
from core.calculator import parse_absorbance_values, calculate_dpph, SAMPLE_EXAMPLE_TEXT
from core.plotter import generate_activity_vs_time_chart, generate_ic50_regression_chart
from core.excel_exporter import export_dpph_excel


class TestDPPHAnalysis(unittest.TestCase):
    def test_parse_absorbance(self):
        text = "1,048\t0.816 \n 0.601;0.539"
        vals = parse_absorbance_values(text)
        self.assertEqual(len(vals), 4)
        self.assertAlmostEqual(vals[0], 1.048)
        self.assertAlmostEqual(vals[1], 0.816)

    def test_calculation(self):
        vals = parse_absorbance_values(SAMPLE_EXAMPLE_TEXT)
        self.assertEqual(len(vals), 25)

        res = calculate_dpph(vals, selected_time=30)
        self.assertIsNotNone(res["ic50"])
        # Expected IC50 approx 41.02
        self.assertAlmostEqual(res["ic50"], 41.0236, places=3)
        self.assertAlmostEqual(res["slope"], 0.7271, places=3)
        self.assertAlmostEqual(res["intercept"], 20.1718, places=3)
        self.assertAlmostEqual(res["r2"], 0.7088, places=3)

    def test_plotter(self):
        vals = parse_absorbance_values(SAMPLE_EXAMPLE_TEXT)
        res = calculate_dpph(vals, selected_time=30)

        buf1 = generate_activity_vs_time_chart(res["rows"], "Test Plant")
        self.assertGreater(len(buf1.getvalue()), 10000)

        buf2 = generate_ic50_regression_chart(
            res["selected_rows"],
            res["slope"],
            res["intercept"],
            res["ic50"],
            30,
            "Test Plant"
        )
        self.assertGreater(len(buf2.getvalue()), 10000)

    def test_excel_export(self):
        vals = parse_absorbance_values(SAMPLE_EXAMPLE_TEXT)
        res = calculate_dpph(vals, selected_time=30)

        os.makedirs("tests/output", exist_ok=True)
        out_file = "tests/output/test_dpph.xlsx"
        export_dpph_excel("Curcuma test", res, "DPPH_template.xlsx", out_file)

        self.assertTrue(os.path.exists(out_file))
        wb = openpyxl.load_workbook(out_file)
        ws = wb.active
        self.assertEqual(ws["A1"].value, "Curcuma test")
        self.assertAlmostEqual(float(ws["O5"].value), res["slope"], places=4)
        self.assertAlmostEqual(float(ws["O6"].value), res["intercept"], places=4)
        self.assertAlmostEqual(float(ws["O8"].value), res["ic50"], places=4)


if __name__ == "__main__":
    unittest.main()
